"""Tools the agent can call: scholarly search, DOI lookup, APA formatting,
a per-course source library, and document text extraction."""

from __future__ import annotations

import io
import json
import os
import re
import urllib.parse
import urllib.request
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

USER_AGENT = "msw-agent/0.1 (grad school research assistant)"
HTTP_TIMEOUT = 20


class ToolError(Exception):
    """An error message that should be shown to the model as a failed tool result."""


def _get_json(url: str) -> dict:
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    try:
        with urllib.request.urlopen(req, timeout=HTTP_TIMEOUT) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        if e.code == 404:
            raise ToolError(f"Not found: {url}") from e
        raise ToolError(f"HTTP {e.code} from {urllib.parse.urlparse(url).netloc}") from e
    except urllib.error.URLError as e:
        raise ToolError(f"Could not reach {urllib.parse.urlparse(url).netloc}: {e.reason}") from e


def normalize_doi(doi: str) -> str:
    doi = doi.strip()
    doi = re.sub(r"^(https?://)?(dx\.)?doi\.org/", "", doi, flags=re.I)
    doi = re.sub(r"^doi:\s*", "", doi, flags=re.I)
    if not re.match(r"^10\.\d{4,9}/\S+$", doi):
        raise ToolError(f"'{doi}' does not look like a DOI (expected something like 10.1037/abc123).")
    return doi


# ---------------------------------------------------------------------------
# APA 7 formatting
# ---------------------------------------------------------------------------

def _initials(given: str) -> str:
    parts = re.split(r"[\s]+", given.strip())
    out = []
    for p in parts:
        if not p:
            continue
        # Hyphenated given names keep the hyphen: Jean-Paul -> J.-P.
        out.append("-".join(f"{s[0].upper()}." for s in p.split("-") if s))
    return " ".join(out)


def format_author(family: str, given: str = "") -> str:
    family = family.strip()
    return f"{family}, {_initials(given)}" if given.strip() else family


def format_author_list(authors: list[str]) -> str:
    """APA 7: list up to 20 authors; with 21+, first 19, ellipsis, then the last."""
    if not authors:
        return ""
    if len(authors) == 1:
        return authors[0]
    if len(authors) <= 20:
        return ", ".join(authors[:-1]) + ", & " + authors[-1]
    return ", ".join(authors[:19]) + ", . . . " + authors[-1]


@dataclass
class Work:
    title: str
    authors: list[str]  # already formatted "Family, G. N."
    year: int | None
    container: str = ""  # journal or book title
    volume: str = ""
    issue: str = ""
    pages: str = ""
    doi: str = ""
    url: str = ""
    kind: str = "article"
    publisher: str = ""
    warnings: list[str] = field(default_factory=list)

    def apa(self) -> str:
        """Plain-text APA 7 reference. Italics are marked with *asterisks*."""
        author_part = format_author_list(self.authors)
        if author_part and not author_part.endswith("."):
            author_part += "."  # group authors: "National Association of Social Workers."
        year_part = f"({self.year})." if self.year else "(n.d.)."
        title = self.title.rstrip(".")
        link = f"https://doi.org/{self.doi}" if self.doi else self.url

        if self.kind == "book":
            head = f"*{title}*."
            tail = f" {self.publisher}." if self.publisher else ""
            # APA: with no author, the title moves to the author position.
            ref = f"{author_part} {year_part} {head}{tail}" if author_part else f"{head} {year_part}{tail}"
        else:
            ref = f"{author_part} {year_part} {title}." if author_part else f"{title}. {year_part}"
            if self.container:
                ref += f" *{self.container}*"
                if self.volume:
                    ref += f", *{self.volume}*"
                    if self.issue:
                        ref += f"({self.issue})"
                if self.pages:
                    ref += f", {self.pages.replace('-', '–')}"
                ref += "."
        if link:
            ref += f" {link}"
        return ref.strip()


_CROSSREF_BOOK_TYPES = {"book", "monograph", "edited-book", "reference-book"}


def work_from_crossref(msg: dict) -> Work:
    warnings = []
    authors = []
    for a in msg.get("author", []) or []:
        if a.get("family"):
            authors.append(format_author(a["family"], a.get("given", "")))
        elif a.get("name"):  # group author, e.g. "National Association of Social Workers"
            authors.append(a["name"])
    if not authors:
        warnings.append("No authors listed in Crossref; check the original source.")

    year = None
    for key in ("published-print", "published-online", "issued", "created"):
        parts = (msg.get(key) or {}).get("date-parts") or []
        if parts and parts[0] and parts[0][0]:
            year = int(parts[0][0])
            break

    title = " ".join(msg.get("title") or []).strip()
    subtitle = " ".join(msg.get("subtitle") or []).strip()
    if subtitle and subtitle.lower() not in title.lower():
        title = f"{title}: {subtitle}"
    title = re.sub(r"<[^>]+>", "", title)  # Crossref titles sometimes contain markup
    if title and title == title.title() and len(title.split()) > 3:
        warnings.append(
            "Title appears to be in Title Case; APA 7 uses sentence case for article and "
            "book titles (capitalize only the first word, first word after a colon, and proper nouns)."
        )

    kind = "book" if msg.get("type") in _CROSSREF_BOOK_TYPES else "article"
    return Work(
        title=title or "[Untitled]",
        authors=authors,
        year=year,
        container=" ".join(msg.get("container-title") or []),
        volume=str(msg.get("volume") or ""),
        issue=str(msg.get("issue") or ""),
        pages=str(msg.get("page") or ""),
        doi=msg.get("DOI", ""),
        url=msg.get("URL", ""),
        kind=kind,
        publisher=msg.get("publisher", "") if kind == "book" else "",
        warnings=warnings,
    )


def fetch_crossref(doi: str) -> Work:
    doi = normalize_doi(doi)
    data = _get_json("https://api.crossref.org/works/" + urllib.parse.quote(doi, safe="/"))
    return work_from_crossref(data.get("message", {}))


# ---------------------------------------------------------------------------
# OpenAlex search
# ---------------------------------------------------------------------------

def _abstract_from_index(inv: dict | None, max_words: int = 180) -> str:
    if not inv:
        return ""
    positions = [(pos, word) for word, poss in inv.items() for pos in poss]
    words = [w for _, w in sorted(positions)]
    text = " ".join(words[:max_words])
    return text + (" …" if len(words) > max_words else "")


def search_openalex(
    query: str,
    from_year: int | None = None,
    to_year: int | None = None,
    articles_only: bool = True,
    limit: int = 8,
) -> list[dict]:
    limit = max(1, min(int(limit), 20))
    filters = []
    if from_year:
        filters.append(f"from_publication_date:{int(from_year)}-01-01")
    if to_year:
        filters.append(f"to_publication_date:{int(to_year)}-12-31")
    if articles_only:
        filters.append("type:article")
    params = {
        "search": query,
        "per_page": str(limit),
        "select": "id,doi,title,publication_year,authorships,primary_location,"
        "cited_by_count,abstract_inverted_index,open_access,type",
    }
    if filters:
        params["filter"] = ",".join(filters)
    if os.environ.get("OPENALEX_API_KEY"):
        params["api_key"] = os.environ["OPENALEX_API_KEY"]
    if os.environ.get("OPENALEX_EMAIL"):
        params["mailto"] = os.environ["OPENALEX_EMAIL"]

    data = _get_json("https://api.openalex.org/works?" + urllib.parse.urlencode(params))
    results = []
    for w in data.get("results", []):
        loc = w.get("primary_location") or {}
        source = loc.get("source") or {}
        authors = [
            (a.get("author") or {}).get("display_name", "")
            for a in (w.get("authorships") or [])
        ]
        results.append(
            {
                "title": w.get("title"),
                "year": w.get("publication_year"),
                "authors": authors[:6] + (["et al."] if len(authors) > 6 else []),
                "venue": source.get("display_name"),
                "doi": (w.get("doi") or "").replace("https://doi.org/", "") or None,
                "type": w.get("type"),
                "cited_by": w.get("cited_by_count"),
                "open_access_url": (w.get("open_access") or {}).get("oa_url"),
                "abstract": _abstract_from_index(w.get("abstract_inverted_index")),
            }
        )
    return results


# ---------------------------------------------------------------------------
# Source library (JSON file on disk, grouped by course)
# ---------------------------------------------------------------------------

class Library:
    def __init__(self, path: Path):
        self.path = Path(path)

    def _load(self) -> dict:
        if not self.path.exists():
            return {"courses": {}}
        return json.loads(self.path.read_text(encoding="utf-8"))

    def _save(self, data: dict) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.path.with_suffix(".tmp")
        tmp.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
        tmp.replace(self.path)

    def add(self, course: str, entry: dict) -> bool:
        """Add or update an entry (keyed by DOI). Returns True if it was new."""
        data = self._load()
        items = data["courses"].setdefault(course.strip() or "General", [])
        for i, existing in enumerate(items):
            if existing.get("doi") and existing["doi"].lower() == entry.get("doi", "").lower():
                items[i] = {**existing, **entry}
                self._save(data)
                return False
        items.append(entry)
        self._save(data)
        return True

    def remove(self, course: str, doi: str) -> bool:
        data = self._load()
        items = data["courses"].get(course, [])
        kept = [e for e in items if e.get("doi", "").lower() != doi.lower()]
        if len(kept) == len(items):
            return False
        data["courses"][course] = kept
        self._save(data)
        return True

    def courses(self) -> list[str]:
        return sorted(self._load()["courses"].keys())

    def entries(self, course: str | None = None) -> list[dict]:
        courses = self._load()["courses"]
        if course:
            return list(courses.get(course, []))
        return [e for items in courses.values() for e in items]

    def reference_list(self, course: str) -> str:
        """APA reference list, alphabetized by first author."""
        refs = sorted((e["apa"] for e in self.entries(course)), key=str.lower)
        return "\n\n".join(refs)


# ---------------------------------------------------------------------------
# Document text extraction (drafts, rubrics, assignment prompts)
# ---------------------------------------------------------------------------

def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages).strip()
    if name.endswith(".docx"):
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:  # rubrics are often tables
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts).strip()
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="replace").strip()
    raise ValueError(f"Unsupported file type: {filename} (use .docx, .pdf, .txt or .md)")


# ---------------------------------------------------------------------------
# Tool definitions exposed to Claude
# ---------------------------------------------------------------------------

WEB_SEARCH = {"type": "web_search_20260209", "name": "web_search", "max_uses": 8}
WEB_FETCH = {"type": "web_fetch_20260209", "name": "web_fetch", "max_uses": 8}

SEARCH_LITERATURE = {
    "name": "search_literature",
    "description": (
        "Search OpenAlex (a free index of ~250M scholarly works) for journal articles and "
        "other research. Returns title, authors, year, journal, DOI, citation count, an "
        "open-access link when one exists, and the start of the abstract. Use this first for "
        "any request for peer-reviewed sources. Results are only as good as the query: try "
        "specific terms (population, intervention, outcome) and run several searches."
    ),
    "input_schema": {
        "type": "object",
        "properties": {
            "query": {"type": "string", "description": "Search terms."},
            "from_year": {"type": "integer", "description": "Earliest publication year."},
            "to_year": {"type": "integer", "description": "Latest publication year."},
            "articles_only": {
                "type": "boolean",
                "description": "Only journal articles (default true). Set false to include books, reports, dissertations.",
            },
            "limit": {"type": "integer", "description": "Number of results, 1-20 (default 8)."},
        },
        "required": ["query"],
    },
}

GET_CITATION = {
    "name": "get_citation",
    "description": (
        "Look up a DOI in Crossref and return verified bibliographic metadata plus an APA 7 "
        "reference built from it. Use this for every source you cite, and to check whether a "
        "reference in the student's paper is real and correctly cited. Any 'warnings' in the "
        "result must be passed on to the student."
    ),
    "input_schema": {
        "type": "object",
        "properties": {"doi": {"type": "string", "description": "DOI, e.g. 10.1093/sw/swx001"}},
        "required": ["doi"],
    },
}

SAVE_SOURCE = {
    "name": "save_source",
    "description": (
        "Save a source to the student's library for a course, with the verified APA reference "
        "and your notes. Only save sources the student asked to keep, or when they asked you to "
        "build a source list. The DOI is looked up in Crossref; sources without a DOI cannot be saved."
    ),
    "input_schema": {
        "type": "object",
        "properties": {
            "course": {"type": "string", "description": "Course name or code, e.g. 'SW 601 Policy'."},
            "doi": {"type": "string"},
            "summary": {"type": "string", "description": "2-4 sentence summary: design, sample, main findings."},
            "relevance": {"type": "string", "description": "How it supports the student's assignment."},
        },
        "required": ["course", "doi", "summary"],
    },
}

LIST_SOURCES = {
    "name": "list_sources",
    "description": "List the sources saved in the student's library, optionally for one course.",
    "input_schema": {
        "type": "object",
        "properties": {"course": {"type": "string", "description": "Omit to list all courses."}},
    },
}


class ToolBox:
    """Holds the tool implementations and dispatches calls from the model."""

    def __init__(self, library: Library):
        self.library = library
        self._handlers: dict[str, Callable[[dict], Any]] = {
            "search_literature": self._search_literature,
            "get_citation": self._get_citation,
            "save_source": self._save_source,
            "list_sources": self._list_sources,
        }

    def run(self, name: str, tool_input: dict) -> tuple[str, bool]:
        """Returns (result_text, is_error)."""
        handler = self._handlers.get(name)
        if handler is None:
            return f"Unknown tool: {name}", True
        try:
            return json.dumps(handler(tool_input), ensure_ascii=False, indent=1), False
        except ToolError as e:
            return str(e), True
        except (KeyError, TypeError, ValueError) as e:
            return f"Invalid input for {name}: {e}", True

    def _search_literature(self, inp: dict) -> Any:
        results = search_openalex(
            inp["query"],
            from_year=inp.get("from_year"),
            to_year=inp.get("to_year"),
            articles_only=inp.get("articles_only", True),
            limit=inp.get("limit", 8),
        )
        return results or {"results": [], "note": "No matches. Try broader or different terms."}

    def _get_citation(self, inp: dict) -> Any:
        work = fetch_crossref(inp["doi"])
        return {
            "apa_reference": work.apa(),
            "title": work.title,
            "authors": work.authors,
            "year": work.year,
            "journal_or_publisher": work.container or work.publisher,
            "doi": work.doi,
            "warnings": work.warnings,
        }

    def _save_source(self, inp: dict) -> Any:
        work = fetch_crossref(inp["doi"])
        entry = {
            "doi": work.doi,
            "apa": work.apa(),
            "title": work.title,
            "year": work.year,
            "summary": inp["summary"],
            "relevance": inp.get("relevance", ""),
            "warnings": work.warnings,
        }
        course = inp["course"].strip() or "General"
        is_new = self.library.add(course, entry)
        return {"saved": True, "course": course, "updated_existing": not is_new, "apa": entry["apa"]}

    def _list_sources(self, inp: dict) -> Any:
        course = inp.get("course")
        entries = self.library.entries(course)
        if not entries:
            return {"sources": [], "courses": self.library.courses()}
        return [{k: e.get(k) for k in ("apa", "summary", "relevance")} for e in entries]
