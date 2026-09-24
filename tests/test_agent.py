import io
import json
from types import SimpleNamespace as NS

import pytest

from msw_agent import tools
from msw_agent.agent import FALLBACK_BETA, Agent
from msw_agent.tools import Library, ToolError, Work, extract_text, normalize_doi, work_from_crossref

CROSSREF_ARTICLE = {
    "type": "journal-article",
    "title": ["Trauma-informed care in schools: A systematic review"],
    "author": [
        {"given": "Jane Ann", "family": "Doe"},
        {"given": "Jean-Paul", "family": "Smith"},
        {"name": "National Association of Social Workers"},
    ],
    "issued": {"date-parts": [[2021, 3]]},
    "container-title": ["Social Work"],
    "volume": "66",
    "issue": "2",
    "page": "101-112",
    "DOI": "10.1093/sw/swab001",
}


def test_apa_article():
    w = work_from_crossref(CROSSREF_ARTICLE)
    assert w.apa() == (
        "Doe, J. A., Smith, J.-P., & National Association of Social Workers. (2021). "
        "Trauma-informed care in schools: A systematic review. *Social Work*, *66*(2), 101–112. "
        "https://doi.org/10.1093/sw/swab001"
    )
    assert w.warnings == []


def test_apa_book_and_title_case_warning():
    w = work_from_crossref(
        {
            "type": "book",
            "title": ["The Strengths Perspective In Social Work Practice"],
            "author": [{"given": "Dennis", "family": "Saleebey"}],
            "issued": {"date-parts": [[2012]]},
            "publisher": "Pearson",
            "DOI": "10.0000/example",
        }
    )
    assert w.apa().startswith("Saleebey, D. (2012). *The Strengths Perspective In Social Work Practice*. Pearson.")
    assert any("sentence case" in x for x in w.warnings)


def test_apa_no_author_moves_title_first():
    w = Work(title="Code of ethics", authors=[], year=2021, kind="book", publisher="NASW Press")
    assert w.apa() == "*Code of ethics*. (2021). NASW Press."


def test_author_list_over_20():
    names = [f"A{i}, B." for i in range(25)]
    out = tools.format_author_list(names)
    assert out.endswith(", . . . A24, B.")
    assert "A19, B." not in out


@pytest.mark.parametrize(
    "raw", ["10.1093/sw/swab001", "https://doi.org/10.1093/sw/swab001", "doi: 10.1093/sw/swab001"]
)
def test_normalize_doi(raw):
    assert normalize_doi(raw) == "10.1093/sw/swab001"


def test_normalize_doi_rejects_garbage():
    with pytest.raises(ToolError):
        normalize_doi("not a doi")


def test_library_roundtrip(tmp_path):
    lib = Library(tmp_path / "lib.json")
    assert lib.add("SW 601", {"doi": "10.1/b", "apa": "Zed, A. (2020). B.", "title": "B"})
    assert lib.add("SW 601", {"doi": "10.1/a", "apa": "Abe, A. (2020). A.", "title": "A"})
    assert not lib.add("SW 601", {"doi": "10.1/A", "apa": "Abe, A. (2020). A.", "title": "A2"})
    assert lib.courses() == ["SW 601"]
    assert lib.reference_list("SW 601") == "Abe, A. (2020). A.\n\nZed, A. (2020). B."
    assert lib.remove("SW 601", "10.1/b")
    assert len(lib.entries("SW 601")) == 1


def test_extract_docx_includes_tables():
    import docx

    d = docx.Document()
    d.add_paragraph("My draft paragraph.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Thesis"
    t.rows[0].cells[1].text = "10 pts"
    buf = io.BytesIO()
    d.save(buf)
    text = extract_text("draft.docx", buf.getvalue())
    assert "My draft paragraph." in text and "Thesis | 10 pts" in text


def test_openalex_abstract_reconstruction():
    assert tools._abstract_from_index({"care": [1], "Trauma": [0], "matters": [2]}) == "Trauma care matters"


# --- Agent loop with a fake API client ---------------------------------------

def block(type, **kw):
    return NS(type=type, **kw)


def resp(stop_reason, *content):
    usage = NS(input_tokens=10, output_tokens=5, cache_read_input_tokens=0)
    return NS(stop_reason=stop_reason, content=list(content), usage=usage)


class FakeClient:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []
        self.beta = NS(messages=NS(create=self._create))

    def _create(self, **kwargs):
        # Snapshot the messages: the agent mutates the same list afterwards.
        self.calls.append({**kwargs, "messages": list(kwargs["messages"])})
        return self.responses.pop(0)


def test_agent_runs_tools_and_returns_answer(tmp_path, monkeypatch):
    monkeypatch.setattr(tools, "fetch_crossref", lambda doi: work_from_crossref(CROSSREF_ARTICLE))
    client = FakeClient(
        [
            resp("tool_use", block("tool_use", id="t1", name="get_citation", input={"doi": "10.1093/sw/swab001"})),
            resp("end_turn", block("text", text="Here is your source.")),
        ]
    )
    agent = Agent(mode="research", client=client, library=Library(tmp_path / "l.json"), model="claude-opus-5")
    events = []
    reply = agent.send("find sources", on_event=lambda k, d: events.append((k, d)))

    assert reply.text == "Here is your source."
    assert events == [("tool", "get_citation: 10.1093/sw/swab001")]
    tool_msg = client.calls[1]["messages"][-1]
    result = tool_msg["content"][0]
    assert result["tool_use_id"] == "t1" and not result["is_error"]
    assert "Doe, J. A." in json.loads(result["content"])["apa_reference"]
    assert client.calls[0]["fallbacks"] == "default" and client.calls[0]["betas"] == [FALLBACK_BETA]


def test_agent_reports_tool_errors(tmp_path):
    client = FakeClient(
        [
            resp("tool_use", block("tool_use", id="t1", name="get_citation", input={"doi": "garbage"})),
            resp("end_turn", block("text", text="That DOI is invalid.")),
        ]
    )
    agent = Agent(mode="feedback", client=client, library=Library(tmp_path / "l.json"))
    agent.send("check")
    result = client.calls[1]["messages"][-1]["content"][0]
    assert result["is_error"] is True


def test_agent_resumes_pause_turn(tmp_path):
    client = FakeClient(
        [
            resp("pause_turn", block("server_tool_use", name="web_search", input={"query": "NASW"})),
            resp("end_turn", block("text", text="Done.")),
        ]
    )
    agent = Agent(client=client, library=Library(tmp_path / "l.json"))
    assert agent.send("q").text == "Done."
    assert client.calls[1]["messages"][-1]["role"] == "assistant"


def test_agent_drops_cut_off_tool_call(tmp_path):
    client = FakeClient(
        [resp("max_tokens", block("text", text="Partial"), block("tool_use", id="t1", name="list_sources", input={}))]
    )
    agent = Agent(client=client, library=Library(tmp_path / "l.json"))
    reply = agent.send("q")
    assert [b.type for b in agent.messages[-1]["content"]] == ["text"]
    assert reply.notices and "length limit" in reply.notices[0]


def test_no_fallbacks_for_other_models(tmp_path):
    client = FakeClient([resp("end_turn", block("text", text="hi"))])
    Agent(client=client, library=Library(tmp_path / "l.json"), model="claude-sonnet-5").send("q")
    assert "fallbacks" not in client.calls[0]
